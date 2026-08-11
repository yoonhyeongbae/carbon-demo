from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

OUT = Path('dist/legal')
OUT.mkdir(parents=True, exist_ok=True)
TODAY = '2026-08-11'
SOFTWARE = 'EV_Carbon_Optimizer'

SOURCES = {
    'sw_act_58': ('소프트웨어 진흥법 제58조', 'https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0058&lsiSeq=265845&urlMode=lsScJoRltInfoR'),
    'sw_rule_17': ('소프트웨어 진흥법 시행규칙 제17조', 'https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0017&lsiSeq=271031&urlMode=lsScJoRltInfoR'),
    'copyright_53': ('저작권법 제53조', 'https://law.go.kr/LSW/lsLawLinkInfo.do?chrClsCd=010202&lsJoLnkSeq=1000979095'),
    'copyright_20': ('저작권법 제20조', 'https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0020&lsiSeq=283335&urlMode=lsScJoRltInfoR'),
    'copyright_rule_6': ('저작권법 시행규칙 제6조', 'https://www.law.go.kr/LSW/lsLinkCommonInfo.do?lsJoLnkSeq=1021643543'),
    'kcc_forms': ('한국저작권위원회 프로그램등록 서식다운로드', 'https://www.copyright.or.kr/customer-center/download-service/customer-support-form/index.do'),
    'natural_terms': ('Natural Earth Terms of Use', 'https://www.naturalearthdata.com/about/terms-of-use/'),
    'natural_countries': ('Natural Earth 1:110m Admin 0 – Countries', 'https://www.naturalearthdata.com/downloads/110m-cultural-vectors/110m-admin-0-countries/'),
    'flet_map': ('Flet Map 공식 문서', 'https://flet.dev/docs/controls/map/'),
    'ortools': ('Google OR-Tools 공식 저장소/Apache-2.0', 'https://github.com/google/or-tools'),
    'vat_8': ('부가가치세법 제8조', 'https://www.law.go.kr/LSW/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1031738661'),
    'vat_decree_11': ('부가가치세법 시행령 제11조', 'https://www.law.go.kr/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1028458807'),
    'ecommerce_12_13': ('전자상거래법 제12조·제13조', 'https://law.go.kr/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1022341999'),
    'ecommerce_rule_8': ('전자상거래법 시행규칙 제8조', 'https://www.law.go.kr/LSW/lsInfoP.do?lsiSeq=166284'),
    'mail_exemption': ('통신판매업 신고 면제 기준에 대한 고시', 'https://www.law.go.kr/admRulLsInfoP.do?admRulSeq=2100000171369'),
}


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_doc_defaults(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Mm(18)
    sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(20)
    sec.right_margin = Mm(20)
    styles = doc.styles
    styles['Normal'].font.name = 'Malgun Gothic'
    styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    styles['Normal'].font.size = Pt(10)
    for name, size, color in [('Title', 22, '1F4E79'), ('Heading 1', 16, '1F4E79'), ('Heading 2', 13, '2F5597'), ('Heading 3', 11, '365F91')]:
        s = styles[name]
        s.font.name = 'Malgun Gothic'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)


def add_title(doc: Document, title: str, subtitle: str | None = None):
    p = doc.add_paragraph()
    p.style = doc.styles['Title']
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(subtitle)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(90, 90, 90)


def add_note(doc: Document, text: str, kind: str = 'info'):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, 'EAF2F8' if kind == 'info' else 'FFF2CC')
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    if kind == 'warn':
        r.bold = True


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        set_cell_shading(c, 'D9EAF7')
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(h)
        rr.bold = True
        rr.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.add_run(str(value)).font.size = Pt(8.7)
    return table


def add_sources(doc: Document, keys: list[str]):
    doc.add_heading('공식 참고문헌', level=2)
    for key in keys:
        title, url = SOURCES[key]
        p = doc.add_paragraph()
        p.add_run(title + ': ').bold = True
        p.add_run(url)


def create_master_doc():
    doc = Document()
    set_doc_defaults(doc)
    add_title(doc, f'{SOFTWARE} 소프트웨어 등록·저작권 등록·배포·상업판매 법적 검토 및 실제 제출 준비서', f'검토 기준일: {TODAY} · 대한민국 법령 기준')
    add_note(doc, '이 문서는 실제 신청 준비에 사용할 수 있도록 작성한 체크리스트·초안입니다. 다만 신청인의 법적 형태(개인/법인/산학협력단), 권리귀속, 판매방식(B2B/B2C), 사업장·주소·대표자 등은 확인되지 않았으므로 빈칸 또는 확인 필요로 두었습니다. 공식 신청서 자체는 제출일에 해당 기관의 최신 서식을 다시 내려받아 사용해야 합니다.', 'warn')

    doc.add_heading('0. 핵심 결론', level=1)
    add_table(doc, ['구분', '현재 소프트웨어의 가능 여부', '핵심 전제', '정부 제출 성격'], [
        ['① 소프트웨어 등록', '가능(단, 정확한 제도명은 소프트웨어사업자 일반 현황 관리)', '신청주체가 소프트웨어사업자이고 사업자 정보가 확정되어야 함', '소프트웨어 제품 자체의 의무등록이 아니라 사업자 현황/실적 관리 신청'],
        ['② 프로그램 저작권 등록', '원칙적으로 가능', '저작자·저작재산권자와 인간의 창작적 기여를 특정하고 프로그램 복제물을 준비', '한국저작권위원회 프로그램등록'],
        ['③ 소프트웨어 배포', '원칙적으로 가능', '배포권한, 제3자 라이선스, 지도·오픈소스 조건, 릴리스 문서 정리', '일반 데스크톱 SW에는 보편적 사전 배포허가/신고 없음'],
        ['④ 상업적 판매', '원칙적으로 가능', '권리귀속 + 사업자등록 + 판매방식에 따른 통신판매/소비자보호 요건', '사업자등록은 기본, B2C 온라인 판매는 통신판매업 신고가 원칙(면제기준 확인)'],
    ])

    doc.add_heading('1. 공급망 지도 라이선스 변경', level=1)
    doc.add_paragraph('CARTO/외부 OpenStreetMap 타일 서비스 의존을 제거하고 Natural Earth 1:110m Admin-0 국가경계 벡터 데이터를 소프트웨어에 로컬로 포함하는 구조를 권장·적용 대상으로 정리한다.')
    add_bullets(doc, [
        'Natural Earth 공식 이용약관은 사이트에서 제공되는 raster/vector 데이터를 public domain으로 밝히고, 수정·전자적 배포·상업적 이용을 허용하며 별도 허가와 저작자 표시를 요구하지 않는다.',
        '따라서 외부 타일 공급자의 상업적 이용조건·rate limit·attribution 정책에 의존하는 위험을 크게 줄일 수 있다.',
        '소프트웨어 화면에는 법적 의무는 아니지만 투명성을 위해 “Made with Natural Earth”를 표시하는 방식을 권장한다.',
        'Natural Earth는 de facto 경계를 기본으로 사용하므로 국가경계의 정치적 표현 문제는 라이선스와 별개로 검토해야 한다.',
        'Flet Map 자체는 PolygonLayer를 제공하므로 로컬 국가 폴리곤을 렌더링할 수 있다. OR-Tools는 Apache License 2.0이다.'
    ])
    add_sources(doc, ['natural_terms', 'natural_countries', 'flet_map', 'ortools'])

    doc.add_heading('2. ① 소프트웨어 등록: 소프트웨어사업자 일반 현황 관리', level=1)
    add_note(doc, '중요: 일반 민간 소프트웨어 제품에 대하여 “제품 자체를 정부에 의무 등록”하는 보편적 절차로 이해하면 부정확합니다. 여기서는 사용자가 요청한 “소프트웨어 등록”을 소프트웨어 진흥법 제58조의 소프트웨어사업자 실적 등 관리 및 시행규칙 제17조의 일반 현황 관리신청으로 정리합니다.', 'warn')
    doc.add_heading('2.1 법적 근거 및 해당성', level=2)
    doc.add_paragraph('소프트웨어 진흥법 제58조는 소프트웨어사업자의 기술인력·사업수행 실적 자료의 제출 및 유지·관리 근거를 두고 있다. 시행규칙 제17조는 별지 제25호서식 “소프트웨어사업자 일반 현황 관리신청서”와 첨부자료를 규정한다. EV_Carbon_Optimizer는 소프트웨어 개발·공급 사업의 대상물이 될 수 있으므로 신청주체가 소프트웨어사업자로 정리되면 이 절차를 이용할 수 있다.')
    doc.add_heading('2.2 실제 준비서류', level=2)
    add_table(doc, ['서류', '필수/조건부', '현재 준비상태', '작성/확보 방법'], [
        ['별지 제25호서식 소프트웨어사업자 일반 현황 관리신청서', '기본', '미작성', '제출 직전 국가법령정보센터/사업자실적관리기관 최신 서식 사용'],
        ['사업자등록증', '행정정보 확인 또는 첨부', '신청주체 미확정', '개인/법인/산학협력단 중 실제 사업주체 확정'],
        ['법인 등기사항증명서', '법인인 경우 행정정보 확인 또는 첨부', '조건부', '법인 신청 시'],
        ['휴·폐업사실증명', '행정정보 확인', '해당 시', '기관 확인'],
        ['최근 연도 결산 재무제표증명', '기관 확인 대상', '사업주체에 따라 준비', '세무자료'],
        ['부가가치세과세표준증명원', '개인사업자 해당', '조건부', '세무자료'],
        ['중소기업확인서', '중소기업 해당', '조건부', '중소기업현황정보시스템 등'],
        ['중견기업확인서', '중견기업 해당', '조건부', '해당 시 첨부'],
        ['실적/변경사항 객관적 증빙', '실적을 기재하는 경우', '향후 가능', '계약서, 납품/검수/세금계산서 등'],
    ])
    doc.add_heading('2.3 절차', level=2)
    add_numbered(doc, ['소프트웨어사업의 실제 신청주체와 사업자등록 상태를 확정한다.', '최신 별지 제25호서식을 내려받아 일반 현황을 작성한다.', '행정정보 공동이용 동의 여부를 정하고, 미동의 항목은 증명서를 직접 첨부한다.', '중소/중견기업 해당 여부와 최근 재무자료를 확인한다.', '사업자실적관리기관에 신청하고, 필요한 경우 별지 제26호 실적관리신청을 별도 진행한다.', '필요 시 별지 제27호 일반 현황 관리확인서를 발급받아 보관한다.'])
    add_sources(doc, ['sw_act_58', 'sw_rule_17'])

    doc.add_heading('3. ② 프로그램 저작권 등록', level=1)
    doc.add_heading('3.1 법적 근거 및 해당성', level=2)
    doc.add_paragraph('저작권법 제53조는 저작자·저작물 제호·종류·창작연월일·공표사항 등의 등록 근거와 일정한 추정효를 규정한다. 저작권법 시행규칙 제6조는 컴퓨터프로그램저작물에 별지 제3호의2 “프로그램등록신청서”와 별지 제4호의2 “프로그램등록신청명세서”를 사용하도록 한다.')
    add_note(doc, 'AI 보조 개발이 포함되었다면 AI를 저작자로 기재하지 말고, 사람이 수행한 요구사항 결정, 공급망/수학모형 설계, 코드 선택·수정·통합, UI/데이터구조 설계, 테스트·검증 등 인간의 창작적 기여와 실제 권리귀속을 증빙하는 편이 안전합니다.', 'warn')
    doc.add_heading('3.2 실제 제출서류', level=2)
    add_table(doc, ['서류/자료', '법적 근거', 'EVCO 준비내용'], [
        ['프로그램등록신청서(별지 제3호의2)', '시행규칙 제6조', '제호: EV_Carbon_Optimizer, 프로그램 종류/저작자/권리자 등 실제 정보 입력'],
        ['프로그램등록신청명세서(별지 제4호의2)', '시행규칙 제6조', '기능, 개발언어, 운영체제, 주요 구성과 창작특징을 기재'],
        ['프로그램 복제물을 수록한 전자적 기록매체 1부', '시행규칙 제6조', '배포 또는 등록 기준 버전의 소스/프로그램 복제물 준비'],
        ['일부 발췌 소스(선택)', '시행규칙 제6조', '일부만으로 창작사실을 입증할 수 있으면 원시 프로그램 언어 형태로 발췌 가능'],
        ['등록사유 증명서류', '증명이 필요한 경우', '창작일/공표일/권리취득 등을 주장할 때 객관적 증빙'],
        ['공동저작자/공동권리자 목록', '공동인 경우', '실제 공동개발 여부에 따라'],
        ['권리자·대리인 증명서류/위임장', '대리신청 또는 권리관계 증명 필요 시', '신청주체와 저작재산권 귀속관계 확정 후 준비'],
        ['인간 창작기여 설명서', '실무 증빙 권고', '요구사항·수학모형·코드·UI·데이터·검증의 인간 기여를 버전관리 기록과 연결'],
        ['버전/해시 목록', '실무 증빙 권고', '제출 프로그램의 SHA-256과 파일 목록 작성'],
    ])
    doc.add_heading('3.3 신청용 서술 초안', level=2)
    doc.add_paragraph('프로그램명: EV_Carbon_Optimizer')
    doc.add_paragraph('프로그램 개요(초안): 프랑스 전기자동차 보조금의 회사 전체 탄소발자국 상한 시나리오를 제약조건으로 두고, 사용자가 설계한 차량 제품·모듈 구조 및 국가·원료·운송수단 공급망 제약을 입력받아 OR-Tools 기반 선형계획으로 총 공급망 비용을 최소화하는 의사결정지원 소프트웨어.')
    doc.add_paragraph('창작적 특징(초안): 동적 차량→부품→재료 및 사용자 정의 모듈 구조, 라인/분산모듈 생산방식의 동적 인덱스, Stage 1→2→3 공급망 물량수지·탄소상한·용량 제약, 시나리오별 결과와 공급망 지도의 통합 시각화.')
    add_sources(doc, ['copyright_53', 'copyright_rule_6', 'kcc_forms'])

    doc.add_heading('4. ③ 소프트웨어 배포', level=1)
    doc.add_heading('4.1 법적 성격', level=2)
    doc.add_paragraph('일반적인 민간 데스크톱 소프트웨어에 대해 배포 전에 정부의 보편적 사전허가를 받는 절차는 확인되지 않는다. 다만 저작권법 제20조에 따라 저작자는 원본·복제물의 배포권을 가지므로, 실제 배포자는 해당 프로그램과 포함 구성요소를 배포할 권리를 확보해야 한다.')
    doc.add_heading('4.2 실제 배포 승인 패키지(내부 제출/보관)', level=2)
    add_table(doc, ['문서', '목적', 'EVCO 적용'], [
        ['배포 승인서/Release Approval', '배포 버전·날짜·책임자 확정', '버전, BUILD ID, SHA-256, 대상 OS 기록'],
        ['제3자 라이선스 목록/SBOM', '재배포 권리 확인', 'Flet, OR-Tools, pandas/numpy 등 실제 의존성 전체 확인'],
        ['THIRD_PARTY_NOTICES', '고지 의무 이행', 'Apache/MIT 등 실제 배포 의존성의 NOTICE/LICENSE 포함 여부 확인'],
        ['Natural Earth 이용근거', '지도 권리 근거', 'Public domain Terms 사본/URL, 데이터 버전 5.1.1 기록'],
        ['EULA 또는 소프트웨어 사용권 계약', '사용범위·복제·역설계·보증·책임 정리', '판매/무상배포 정책에 맞춰 확정'],
        ['개인정보처리방침', '개인정보 처리 시', '계정/로그/클라우드 전송 등 개인정보를 실제 처리하는 경우 작성'],
        ['보안·악성코드·설치 테스트 기록', '배포품 안전성 검증', 'Windows 환경 테스트, 서명/해시 검증'],
        ['릴리스 노트 및 사용자 매뉴얼', '기능/제약/버전 고지', '최적화 가정, 데이터 책임, Scope proxy 설명 포함'],
    ])
    doc.add_heading('4.3 Natural Earth 지도 적용 결론', level=2)
    doc.add_paragraph('Natural Earth 데이터는 공식 이용약관상 public domain이므로 CARTO hosted tiles보다 배포·상업배포 라이선스 관리가 단순하다. EVCO에는 1:110m Admin-0 국가 폴리곤을 로컬 asset으로 번들하고 외부 지도 타일 호출을 제거하는 구조가 적합하다.')
    add_sources(doc, ['copyright_20', 'natural_terms', 'natural_countries', 'flet_map', 'ortools'])

    doc.add_heading('5. ④ 상업적 판매', level=1)
    doc.add_heading('5.1 사업자등록', level=2)
    doc.add_paragraph('부가가치세법 제8조에 따르면 사업자는 원칙적으로 사업 개시일부터 20일 이내 사업자등록을 신청하며, 개시 전에도 신청할 수 있다. 시행령 제11조는 신청서 기재사항과 사업허가/등록/신고가 필요한 업종, 임차·전차 사업장 등 상황별 첨부서류를 정한다.')
    doc.add_heading('5.2 온라인 B2C 판매 시 통신판매', level=2)
    doc.add_paragraph('전자상거래법 제12조는 통신판매업 신고를 원칙으로 하고 일정 소규모 기준 이하에는 예외를 둔다. 시행규칙 제8조는 별지 제1호서식 통신판매업 신고서를 규정한다. 선불식 통신판매 등에는 결제구조에 따라 구매안전서비스 관련 서류가 필요할 수 있다. 판매화면에는 제13조에 따른 사업자 신원 및 거래조건 표시가 필요하다.')
    add_table(doc, ['실제 준비서류', '필수/조건부', '비고'], [
        ['사업자등록 신청', '판매 사업 개시 시 기본', '개인/법인/산학협력단 등 실제 판매주체 확정'],
        ['사업장 임대차계약서 등', '사업장 임차/전차 시', '시행령 제11조 상황별'],
        ['법령상 허가·등록·신고 증명', '해당 업종인 경우', '일반 SW 판매 자체에 별도 허가가 필요한지 사업형태별 확인'],
        ['통신판매업 신고서(별지 제1호)', '온라인 B2C 통신판매 원칙', '법정 면제기준 해당 여부 별도 확인'],
        ['구매안전서비스 이용 확인증', '해당 결제형태 등 조건부', '신고기관/결제방식에 따라 확인'],
        ['EULA/사용권계약', '강력 권고', '라이선스 범위, 사용자수, 설치수, 금지행위, 업데이트'],
        ['판매 이용약관/거래조건', '온라인 판매 권고/필요', '가격, 세금, 결제, 인도, 업데이트, 책임, 분쟁'],
        ['청약철회·환불정책', 'B2C 온라인 판매 시 중요', '디지털콘텐츠 특례와 제공개시 동의 절차는 실제 판매 UI와 함께 법률 검토'],
        ['사업자 신원 표시', 'B2C 온라인 표시·광고/청약 화면', '상호, 대표자, 주소, 연락처, 전자우편, 신고번호 등'],
        ['개인정보처리방침/수탁자 고지', '개인정보 수집·결제·계정 사용 시', '실제 데이터흐름에 맞춰 작성'],
        ['세금계산서/현금영수증/영수증 프로세스', '거래유형별', '회계·세무처리와 연동'],
    ])
    doc.add_heading('5.3 권장 실제 진행 순서', level=2)
    add_numbered(doc, ['저작자·저작재산권자·판매주체를 먼저 확정한다.', '프로그램 저작권 등록용 기준버전을 동결하고 해시를 기록한다.', '사업자등록을 완료한다.', '온라인 B2C 판매이면 통신판매업 신고 및 결제/구매안전 구조를 확정한다.', 'EULA, 판매약관, 환불정책, 개인정보처리방침을 실제 판매채널에 맞춰 확정한다.', '제3자 라이선스와 Natural Earth public-domain 근거를 배포물에 함께 보관한다.', '코드서명/악성코드 검사/설치 테스트 후 상업 릴리스를 승인한다.'])
    add_sources(doc, ['vat_8', 'vat_decree_11', 'ecommerce_12_13', 'ecommerce_rule_8', 'mail_exemption'])

    doc.add_heading('6. 실제 제출 전에 확정해야 할 신청인 정보', level=1)
    add_table(doc, ['항목', '기입란'], [
        ['신청/판매 주체', '[개인 / 법인 / 산학협력단 / 기타:                ]'],
        ['대표자/저작자 성명', '[                                      ]'],
        ['저작재산권자', '[                                      ]'],
        ['사업자등록번호', '[                                      ]'],
        ['사업장 주소', '[                                      ]'],
        ['연락처/전자우편', '[                                      ]'],
        ['프로그램 최초 창작일', '[                                      ]'],
        ['최초 공표 여부/공표일', '[                                      ]'],
        ['등록 기준 버전/BUILD ID', '[                                      ]'],
        ['SHA-256', '[                                      ]'],
        ['판매방식', '[B2B / B2C / 온라인 다운로드 / 라이선스키 / SaaS / 기타]'],
    ])

    doc.add_heading('7. 최종 확인', level=1)
    add_note(doc, '이 문서로 신청 준비를 할 수 있으나, 실제 공식 양식 제출 시에는 반드시 제출 당일의 최신 법령·기관 서식을 다시 확인하세요. 특히 통신판매업 신고 면제기준, 디지털콘텐츠 청약철회 예외, 개인정보 처리, 대학/산학협력단의 직무발명·저작권 귀속은 신청인 상황에 따라 달라질 수 있습니다.', 'warn')

    path = OUT / 'EV_Carbon_Optimizer_4대절차_법적검토_실제제출준비서_2026-08-11.docx'
    doc.save(path)
    return path


def create_process_doc(filename: str, title: str, purpose: str, rows: list[list[str]], source_keys: list[str], notes: list[str]):
    doc = Document()
    set_doc_defaults(doc)
    add_title(doc, title, f'{SOFTWARE} · 실제 제출 준비용 · {TODAY}')
    add_note(doc, purpose, 'info')
    add_table(doc, ['순서', '제출/보관 서류', '상태', '비고'], rows)
    doc.add_heading('작성·제출 전 확인사항', level=1)
    add_bullets(doc, notes)
    add_sources(doc, source_keys)
    path = OUT / filename
    doc.save(path)
    return path


def pdf_font():
    candidates = [
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/opentype/noto/NotoSansCJKkr-Regular.otf',
        '/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc',
    ]
    for p in candidates:
        if Path(p).exists():
            try:
                pdfmetrics.registerFont(TTFont('Korean', p, subfontIndex=1))
                return 'Korean'
            except Exception:
                try:
                    pdfmetrics.registerFont(TTFont('Korean', p))
                    return 'Korean'
                except Exception:
                    pass
    return 'Helvetica'


def create_reference_pdf(filename: str, title: str, summary: str, refs: list[tuple[str, str, str]]):
    font = pdf_font()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='KTitle', parent=styles['Title'], fontName=font, fontSize=18, leading=24, alignment=TA_CENTER, textColor=colors.HexColor('#1F4E79')))
    styles.add(ParagraphStyle(name='KH1', parent=styles['Heading2'], fontName=font, fontSize=12, leading=17, textColor=colors.HexColor('#2F5597')))
    styles.add(ParagraphStyle(name='KBody', parent=styles['BodyText'], fontName=font, fontSize=9.5, leading=14))
    doc = SimpleDocTemplate(str(OUT / filename), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    story = [Paragraph(title, styles['KTitle']), Spacer(1, 6*mm), Paragraph(summary, styles['KBody']), Spacer(1, 5*mm)]
    for idx, (name, url, applicability) in enumerate(refs, 1):
        story.append(Paragraph(f'{idx}. {name}', styles['KH1']))
        story.append(Paragraph(f'공식 위치: {url}', styles['KBody']))
        story.append(Paragraph(f'EVCO 적용: {applicability}', styles['KBody']))
        story.append(Spacer(1, 4*mm))
    story.append(Spacer(1, 4*mm))
    story.append(Paragraph('검증방법: 위 공식 URL을 직접 열어 시행일·조문·별지서식을 제출 당일 다시 확인하십시오. 본 PDF는 공식 법령 원문을 대체하는 서류가 아니라 검증용 인덱스입니다.', styles['KBody']))
    doc.build(story)
    return OUT / filename


def main():
    files = [create_master_doc()]
    files.append(create_process_doc(
        '01_소프트웨어사업자_일반현황_실제신청_준비서.docx',
        '① 소프트웨어사업자 일반 현황 관리 실제 신청 준비서',
        '일반 민간 소프트웨어 제품 자체의 의무등록이 아니라 소프트웨어사업자의 일반 현황/실적 관리 신청을 위한 준비서입니다.',
        [
            ['1', '별지 제25호 일반 현황 관리신청서', '□', '최신 공식양식 다운로드'],
            ['2', '사업자등록증/법인등기사항 등', '□', '행정정보 공동이용 동의 시 기관 확인 가능 항목 존재'],
            ['3', '최근 결산 재무제표 등', '□', '사업자 유형에 따라'],
            ['4', '중소/중견기업 확인서', '□', '해당 시'],
            ['5', '실적 객관증빙', '□', '실적 신청 시 계약서·검수 등'],
        ],
        ['sw_act_58', 'sw_rule_17'],
        ['신청주체의 법적 형태와 사업자등록 상태를 먼저 확정합니다.', '제품 자체의 등록증이 발급되는 절차와 혼동하지 않습니다.', '실적관리까지 진행하면 별지 제26호 및 실적증빙을 추가합니다.'],
    ))
    files.append(create_process_doc(
        '02_프로그램저작권등록_실제제출_준비서.docx',
        '② EV_Carbon_Optimizer 프로그램 저작권 등록 실제 제출 준비서',
        '한국저작권위원회 프로그램등록 신청에 필요한 법정 서류와 EVCO 증빙을 준비하기 위한 체크리스트입니다.',
        [
            ['1', '프로그램등록신청서(별지 제3호의2)', '□', '최신 KCC/법령 서식'],
            ['2', '프로그램등록신청명세서(별지 제4호의2)', '□', '기능·창작특징·개발환경'],
            ['3', '프로그램 복제물 전자적 기록매체 1부', '□', '기준버전 고정 및 해시 기록'],
            ['4', '등록사유/권리관계 증빙', '□', '필요한 경우'],
            ['5', '공동저작자 목록/위임장 등', '□', '해당 시'],
            ['6', '인간 창작기여 설명서', '□', 'AI 보조개발 부분과 인간의 선택·수정·통합을 분리'],
        ],
        ['copyright_53', 'copyright_rule_6', 'kcc_forms'],
        ['저작자와 저작재산권자를 실제 계약·소속관계에 따라 확정합니다.', '창작일을 등록하려면 1년 경과 여부에 따른 추정효 제한을 확인합니다.', '프로그램 일부 발췌를 제출하는 경우 변환 전 원시 프로그램 언어 형태를 사용합니다.'],
    ))
    files.append(create_process_doc(
        '03_소프트웨어배포_릴리스승인_준비서.docx',
        '③ EV_Carbon_Optimizer 소프트웨어 배포 릴리스 승인 준비서',
        '일반 배포에 보편적 정부 사전허가가 있는 것은 아니므로, 실제 배포권·제3자 라이선스·릴리스 품질을 입증하는 내부 배포 패키지로 사용합니다.',
        [
            ['1', '릴리스 승인서/버전·해시', '□', '배포본 고정'],
            ['2', 'SBOM/제3자 라이선스 목록', '□', 'Flet, OR-Tools 등'],
            ['3', 'THIRD_PARTY_NOTICES/LICENSES', '□', '각 라이선스 요건 반영'],
            ['4', 'Natural Earth public-domain 근거', '□', '버전 5.1.1/Terms 기록'],
            ['5', 'EULA/사용자매뉴얼/릴리스노트', '□', '배포조건·가정 고지'],
            ['6', '설치·악성코드·해시 검증기록', '□', 'Windows 릴리스 품질'],
        ],
        ['copyright_20', 'natural_terms', 'natural_countries', 'flet_map', 'ortools'],
        ['CARTO/외부 OSM hosted tile 의존을 제거하고 Natural Earth 로컬 벡터를 번들하는 구성을 사용합니다.', 'Natural Earth는 public domain이지만 정치적 경계표현/정확성은 별도 검토사항입니다.', '상용 배포 전에 실제 설치파일에 포함되는 모든 Python/Flutter 종속성을 SBOM으로 다시 스캔합니다.'],
    ))
    files.append(create_process_doc(
        '04_상업판매_사업자통신판매_준비서.docx',
        '④ EV_Carbon_Optimizer 상업적 판매 실제 준비서',
        '판매주체의 사업자등록과 온라인 B2C 판매 시 통신판매/소비자보호 요건을 준비하기 위한 체크리스트입니다.',
        [
            ['1', '사업자등록 신청', '□', '사업 개시 전 또는 개시 후 법정기한 내'],
            ['2', '사업장·임대차 등 조건부 첨부', '□', '시행령 제11조 확인'],
            ['3', '통신판매업 신고서(별지 제1호)', '□', '온라인 B2C 및 면제기준 확인'],
            ['4', '구매안전서비스 관련 서류', '□', '결제형태/선불거래 등 해당 시'],
            ['5', 'EULA/판매약관/환불정책', '□', '소비자 판매 UI와 일치'],
            ['6', '사업자 신원 표시·개인정보처리방침', '□', '실제 판매채널과 데이터흐름에 맞춤'],
            ['7', '세금/영수증 프로세스', '□', '세무대리인과 확인 권고'],
        ],
        ['vat_8', 'vat_decree_11', 'ecommerce_12_13', 'ecommerce_rule_8', 'mail_exemption'],
        ['B2B 전용 판매와 B2C 온라인 판매는 적용되는 소비자보호 실무가 다릅니다.', '통신판매업 신고 면제기준은 거래횟수/규모와 최신 고시를 실제 판매 직전에 확인합니다.', '디지털콘텐츠 청약철회 제한을 적용하려면 단순 약관 기재만으로 끝내지 말고 실제 사전고지·동의 UI를 법률검토해야 합니다.'],
    ))

    files.append(create_reference_pdf(
        'REF_01_소프트웨어사업자_일반현황_법적참고문헌.pdf',
        '참고문헌 ① 소프트웨어사업자 일반 현황 관리',
        '“소프트웨어 등록”을 제품 의무등록이 아니라 소프트웨어사업자 일반 현황/실적 관리 제도로 정리한 검증용 참고문헌입니다.',
        [
            (SOURCES['sw_act_58'][0], SOURCES['sw_act_58'][1], '소프트웨어사업자 기술인력·사업수행 실적 등 자료 제출 및 유지·관리의 법률상 근거'),
            (SOURCES['sw_rule_17'][0], SOURCES['sw_rule_17'][1], '별지 제25호 일반 현황 관리신청서, 첨부/행정정보 확인자료, 실적관리 및 확인서 근거'),
        ]
    ))
    files.append(create_reference_pdf(
        'REF_02_프로그램저작권등록_법적참고문헌.pdf',
        '참고문헌 ② 프로그램 저작권 등록',
        'EV_Carbon_Optimizer를 컴퓨터프로그램저작물로 등록할 때 확인해야 할 현재 법령·공식 서식 위치입니다.',
        [
            (SOURCES['copyright_53'][0], SOURCES['copyright_53'][1], '저작자·제호·종류·창작일·공표사항 등록 및 추정효'),
            (SOURCES['copyright_rule_6'][0], SOURCES['copyright_rule_6'][1], '별지 제3호의2/제4호의2 및 프로그램 복제물 등 제출자료'),
            (SOURCES['kcc_forms'][0], SOURCES['kcc_forms'][1], '프로그램등록 최신 신청서와 명세서 공식 다운로드 위치'),
        ]
    ))
    files.append(create_reference_pdf(
        'REF_03_소프트웨어배포_법적라이선스_참고문헌.pdf',
        '참고문헌 ③ 소프트웨어 배포 및 지도 라이선스',
        '일반 배포권과 EVCO 배포물의 지도·오픈소스 라이선스 근거를 확인하기 위한 참고문헌입니다.',
        [
            (SOURCES['copyright_20'][0], SOURCES['copyright_20'][1], '저작자의 원본·복제물 배포권'),
            (SOURCES['natural_terms'][0], SOURCES['natural_terms'][1], 'Natural Earth public domain, 수정·전자배포·상업적 이용 가능, 허가/표시 불필요'),
            (SOURCES['natural_countries'][0], SOURCES['natural_countries'][1], 'EVCO 로컬 배경지도에 사용할 1:110m Admin-0 Countries 5.1.1'),
            (SOURCES['flet_map'][0], SOURCES['flet_map'][1], '지도 레이어 기능 및 타일공급자별 이용정책 준수 경고'),
            (SOURCES['ortools'][0], SOURCES['ortools'][1], '최적화 엔진 OR-Tools Apache-2.0'),
        ]
    ))
    files.append(create_reference_pdf(
        'REF_04_상업판매_법적참고문헌.pdf',
        '참고문헌 ④ 상업적 판매',
        '소프트웨어를 실제 유상 판매할 때의 사업자등록과 온라인 B2C 통신판매 관련 핵심 근거입니다.',
        [
            (SOURCES['vat_8'][0], SOURCES['vat_8'][1], '사업 개시 후 20일 내 사업자등록 원칙 및 사전등록 가능'),
            (SOURCES['vat_decree_11'][0], SOURCES['vat_decree_11'][1], '사업자등록 신청 기재사항·상황별 첨부서류'),
            (SOURCES['ecommerce_12_13'][0], SOURCES['ecommerce_12_13'][1], '통신판매업 신고 및 사업자 신원·거래조건 제공'),
            (SOURCES['ecommerce_rule_8'][0], SOURCES['ecommerce_rule_8'][1], '별지 제1호 통신판매업 신고서 및 구매안전서비스 관련 서류'),
            (SOURCES['mail_exemption'][0], SOURCES['mail_exemption'][1], '소규모 통신판매업 신고 면제 기준'),
        ]
    ))

    links = OUT / 'OFFICIAL_SOURCE_LINKS_2026-08-11.txt'
    with links.open('w', encoding='utf-8') as f:
        f.write(f'{SOFTWARE} official legal/source links — verified basis date {TODAY}\n\n')
        for title, url in SOURCES.values():
            f.write(f'- {title}\n  {url}\n')
    files.append(links)

    zip_path = OUT / 'EV_Carbon_Optimizer_v0.5.8_LEGAL_PACK.zip'
    with ZipFile(zip_path, 'w', ZIP_DEFLATED) as z:
        for p in files:
            z.write(p, p.name)
    print('Generated', len(files), 'files plus', zip_path)

if __name__ == '__main__':
    main()
