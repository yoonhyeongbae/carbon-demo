from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml.ns import qn

OUT=Path('dist/legal'); OUT.mkdir(parents=True, exist_ok=True)
doc=Document()
sec=doc.sections[0]; sec.top_margin=Mm(18); sec.bottom_margin=Mm(18); sec.left_margin=Mm(20); sec.right_margin=Mm(20)
for name in ['Normal','Title','Heading 1','Heading 2']:
    s=doc.styles[name]; s.font.name='Malgun Gothic'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
doc.styles['Normal'].font.size=Pt(10); doc.styles['Title'].font.size=Pt(21); doc.styles['Title'].font.color.rgb=RGBColor(31,78,121)
doc.add_heading('EV_Carbon_Optimizer 법령 최신검증 보충서',0)
doc.add_paragraph('검증 기준일: 2026-08-11. 이 보충서는 법적검토 패키지의 법령 링크/시행일을 제출 직전 다시 확인하기 위한 최신 확인표입니다.')

doc.add_heading('1. 소프트웨어사업자 일반 현황 관리',1)
doc.add_paragraph('소프트웨어 진흥법 제58조 및 소프트웨어 진흥법 시행규칙 제17조를 확인합니다. 시행규칙 제17조는 별지 제25호 소프트웨어사업자 일반 현황 관리신청서와 첨부/행정정보 확인자료를 규정합니다.')
doc.add_paragraph('공식: https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0058&lsiSeq=265845&urlMode=lsScJoRltInfoR')
doc.add_paragraph('공식: https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0017&lsiSeq=271031&urlMode=lsScJoRltInfoR')

doc.add_heading('2. 프로그램 저작권 등록',1)
doc.add_paragraph('2026-05-11 시행 저작권법 제53조 및 저작권법 시행규칙 제6조를 확인합니다. 프로그램은 별지 제3호의2 프로그램등록신청서와 별지 제4호의2 프로그램등록신청명세서를 사용하며, 프로그램 복제물을 수록한 전자적 기록매체 1부가 요구됩니다.')
doc.add_paragraph('공식: https://law.go.kr/LSW/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1029423425')
doc.add_paragraph('공식: https://www.law.go.kr/LSW/lsLawLinkInfo.do?chrClsCd=010202&lsJoLnkSeq=900602048')

doc.add_heading('3. 소프트웨어 배포 및 지도',1)
doc.add_paragraph('Natural Earth 공식 Terms는 사이트의 raster/vector map data를 public domain으로 명시하고 수정, 전자적 배포, 상업적 이용을 허용하며 허가와 attribution을 요구하지 않습니다. EVCO v0.5.8은 CARTO/외부 OSM hosted tile을 제거하고 Natural Earth 1:110m Admin-0 로컬 벡터를 사용하는 방향입니다.')
doc.add_paragraph('공식: https://www.naturalearthdata.com/about/terms-of-use/')
doc.add_paragraph('공식 데이터: https://www.naturalearthdata.com/downloads/110m-cultural-vectors/110m-admin-0-countries/')

doc.add_heading('4. 상업적 판매',1)
doc.add_paragraph('부가가치세법 제8조: 사업자는 원칙적으로 사업 개시일부터 20일 이내 사업자등록을 신청하며 개시 전 신청도 가능합니다.')
doc.add_paragraph('전자상거래법 제12조: 통신판매업자는 원칙적으로 신고하되 공정위 고시 면제기준 이하인 경우 예외입니다. 2022-04-05 시행 공정위고시 제2022-4호의 현재 확인 가능한 면제기준은 (1) 직전년도 통신판매 거래횟수 50회 미만 또는 (2) 부가가치세법상 간이과세자입니다. 과거의 “최근 6개월 20회/1,200만원” 기준은 2018년 고시로서 현재 기준으로 사용하지 않습니다.')
doc.add_paragraph('사업자등록 공식: https://www.law.go.kr/LSW/lsLinkCommonInfo.do?chrClsCd=010202&lsJoLnkSeq=1029624963')
doc.add_paragraph('통신판매업 신고 공식: https://www.law.go.kr/LSW/lsSideInfoP.do?docCls=jo&joBrNo=00&joNo=0012&lsiSeq=282793&urlMode=lsScJoRltInfoR')
doc.add_paragraph('시행규칙 제8조/별지 제1호: https://law.go.kr/LSW/lsLinkCommonInfo.do?chrClsCd=010202&lspttninfSeq=63518')
doc.add_paragraph('현행 신고 면제 고시: https://www.law.go.kr/LSW/admRulLsInfoP.do?admRulSeq=2100000210384')

doc.add_heading('5. 실제 제출 원칙',1)
doc.add_paragraph('본 패키지의 Word/PDF는 준비서와 검증 인덱스입니다. 공식 신청서 자체는 제출 당일 국가법령정보센터·한국저작권위원회·관할 세무서/정부24·관할 지자체 등에서 최신 서식을 다시 내려받아 사용해야 합니다. 신청인 법적 형태, 권리귀속, B2B/B2C 판매형태, 개인정보 처리여부에 따라 추가서류가 달라질 수 있습니다.')

path=OUT/'00_법령_최신검증_보충서_2026-08-11.docx'; doc.save(path); print(path)
