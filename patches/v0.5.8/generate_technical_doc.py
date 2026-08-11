from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('dist/update')
OUT.mkdir(parents=True, exist_ok=True)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Mm(18); sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(19); sec.right_margin = Mm(19)
    for sname in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        s = doc.styles[sname]
        s.font.name = 'Malgun Gothic'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Title'].font.color.rgb = RGBColor(31,78,121)
    doc.styles['Heading 1'].font.color.rgb = RGBColor(31,78,121)
    doc.styles['Heading 2'].font.color.rgb = RGBColor(47,85,151)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'D9EAF7'); r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].paragraphs[0].add_run(str(v)).font.size=Pt(8.7)
    return t


def bullets(doc, items):
    for x in items:
        doc.add_paragraph(x, style='List Bullet')


def code(doc, text):
    p=doc.add_paragraph()
    r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(8.5)
    p.paragraph_format.left_indent=Mm(5)


def main():
    doc=Document(); setup(doc)
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('EV_Carbon_Optimizer v0.5.8\n전체 코드 변경·공급망 지도·라이선스 적용 정리')
    p2=doc.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run('기준일 2026-08-11 · 대상: v0.5.6 DEV_READY → v0.5.8 업데이트').font.size=Pt(9.5)

    doc.add_heading('1. 수정 목적', level=1)
    bullets(doc,[
        '5번 결과 탭에서 시나리오×생산방식 체크박스와 함께 Stage 1, Stage 1→2, Stage 2, Stage 2→3을 사용자가 개별 선택하도록 복원한다.',
        '현재 제품구조에서 활성화된 철강, 알루미늄, 기타 원자재, 희토류, 구리, 플라스틱, 배터리 등 원료·원자재를 지도에서 개별 체크/해제하도록 복원한다.',
        '6번 분석 탭에서 5번 결과 탭과 중복되는 Scope 1+2 / Scope 3 proxy 막대그래프를 제거한다.',
        'CARTO 및 외부 OpenStreetMap hosted tile 의존을 제거하고 Natural Earth 1:110m Admin-0 국가경계 벡터를 소프트웨어 내부 asset으로 사용한다.',
        '지도 확대·축소/이동 시 부하를 줄이기 위해 경로 점 수, 화살표 수, 지도 이벤트와 외부 타일 네트워크 요청을 줄인다.',
        '이번 지도/UI 변경은 OR-Tools의 목적함수, 결정변수, 제약조건 및 18개 CSV 데이터 의미를 변경하지 않는다.'
    ])

    doc.add_heading('2. 지도 데이터 변경', level=1)
    table(doc,['항목','v0.5.6','v0.5.8'],[
        ['배경지도','CARTO hosted tile + OpenStreetMap attribution','Natural Earth Admin-0 1:110m 로컬 벡터 폴리곤'],
        ['네트워크 의존','지도 확대/이동 중 외부 tile 요청','배경 국가경계는 로컬 파일에서 읽음'],
        ['법적 성격','타일 공급자 이용약관·attribution·commercial 정책 확인 필요','Natural Earth 공식 Terms상 public domain; 상업적 이용·전자배포 허용'],
        ['표시','CARTO/OSM attribution','Made with Natural Earth (법적 의무는 아니나 투명성 목적)'],
        ['경로 geometry','Bézier 약 45점, 경로당 화살표 3개','Bézier 약 19점, 경로당 화살표 1개'],
    ])
    doc.add_paragraph('Natural Earth 공식 이용약관: https://www.naturalearthdata.com/about/terms-of-use/')
    doc.add_paragraph('사용 데이터: https://www.naturalearthdata.com/downloads/110m-cultural-vectors/110m-admin-0-countries/')
    doc.add_paragraph('주의: public domain 여부와 국경선의 정치적 표현/정확성 문제는 별개이다. Natural Earth의 기본 Admin-0 데이터는 de facto 경계를 사용하므로 특정 시장에 배포할 때는 경계표현 정책을 별도로 검토한다.')

    doc.add_heading('3. 5번 결과 탭 지도 선택 구조', level=1)
    doc.add_paragraph('공통 지도 컨트롤 순서는 다음과 같다.')
    table(doc,['순서','컨트롤','기본값','적용 범위'],[
        ['1','S1/S2/S3 × 라인/모듈 6개 지도 체크박스','6개 모두 선택','표시할 결과 지도 선택'],
        ['2','공통 지도 범례','활성 원료 전체','색상·Stage·경로선 의미'],
        ['3','Stage 1 / Stage 1→2 / Stage 2 / Stage 2→3','4개 모두 선택','선택된 모든 결과지도에 공통 적용'],
        ['4','원료·원자재 체크박스','현재 제품구조 활성 품목 모두 선택','생산지와 Stage1→2 경로를 품목별 표시/숨김'],
        ['5','인터랙티브 지도','선택 조건을 반영','모든 선택 조합'],
    ])
    doc.add_paragraph('Stage 3 프랑스 시장은 별도 체크박스를 만들지 않고 Stage 2→3을 선택했을 때 완성차 경로와 함께 표시한다.')

    doc.add_heading('4. 코드 구성', level=1)
    table(doc,['파일','역할','v0.5.8 변경'],[
        ['main.py','6개 탭 UI, 입력상태, 결과/분석 화면','v0.5.7 shared Stage/material filter patch 적용; 지도 import를 natural_earth_map으로 변경'],
        ['map_logic.py','Stage 1/1→2/2/2→3 지도 payload·곡선·색상','경로 geometry 44→18 step, 화살표 3→1'],
        ['natural_earth_map.py','v0.5.8 신규 로컬 배경지도 renderer','Natural Earth shapefile → Flet PolygonLayer; TileLayer 없음'],
        ['dynamic_structure_core.py','동적 OR-Tools LP','변경 없음'],
        ['legacy_core.py','기준 수식/계수 로직','변경 없음'],
        ['ev_optimizer_api.py','UI와 solver 연결','변경 없음'],
        ['data/*.csv','현재 작업 데이터 18개','변경 없음'],
        ['baseline_data/*.csv','복구용 baseline 18개','변경 없음'],
    ])

    doc.add_heading('5. v0.5.8 업데이트 적용 순서', level=1)
    for i,x in enumerate([
        '기존 EV_Carbon_Optimizer_v0.5.6_DEV_READY 폴더를 별도 백업한다.',
        'v0.5.8 UPDATE ZIP의 파일을 기존 프로젝트 최상위 폴더(main.py와 START_DEV.bat이 있는 위치)에 덮어쓴다.',
        'APPLY_V058_NATURAL_EARTH_PATCH.bat을 실행한다. 필요한 경우 먼저 START_DEV.bat을 한 번 실행해 Python 환경을 준비한다.',
        '업데이트 프로그램이 v0.5.7 shared filter patch를 먼저 적용하고 Natural Earth 데이터를 공식 S3에서 다운로드하여 assets/natural_earth에 저장한다.',
        '업데이트 완료 후 START_DEV.bat을 실행한다.',
        '5번 결과 탭에서 Stage 4개 및 원료별 체크박스가 나타나는지 확인하고, 지도 하단/범례에서 Natural Earth 사용표시를 확인한다.'
    ],1):
        doc.add_paragraph(f'{i}. {x}')

    doc.add_heading('6. 지도 렉 감소 설계', level=1)
    bullets(doc,[
        '경로 Bézier 점을 약 45개에서 약 19개로 축소한다.',
        '각 경로의 방향 화살표를 3개에서 도착지 부근 1개로 줄인다.',
        '숨긴 Stage/품목의 지도 객체는 먼저 필터링하여 생성하지 않는다.',
        'Natural Earth 배경은 PolygonLayer culling/simplification을 사용하여 현재 뷰포트 중심으로 렌더링한다.',
        'Map keep_alive를 사용 가능한 Flet 버전에서 활성화하여 탭 전환/리빌드 부담을 줄인다.',
        '외부 tile server와 통신하지 않으므로 pan/zoom 중 네트워크 tile 다운로드 지연을 제거한다.'
    ])

    doc.add_heading('7. 법적/상업배포 관점의 판단', level=1)
    doc.add_paragraph('Natural Earth는 공식 Terms에서 웹사이트의 raster/vector map data를 public domain으로 명시하고, 수정·전자배포·출력·개인/교육/상업적 이용을 허용하며 별도 허가나 attribution을 요구하지 않는다. 따라서 CARTO hosted tile 서비스의 별도 이용약관·상업조건에 의존하는 구조보다 EVCO의 등록·배포·상업판매 준비에 적합한 선택이다. 다만 EVCO 전체의 배포 가능성은 OR-Tools, Flet 및 실제 배포물에 포함되는 모든 제3자 구성요소의 라이선스 준수와 별개로 확인해야 한다.')

    doc.add_heading('8. 검증 체크리스트', level=1)
    bullets(doc,[
        '□ 5번 결과 탭: S1/S2/S3 × 라인/모듈 6개 선택 가능',
        '□ 공통 범례가 시나리오 체크박스 바로 아래에 표시',
        '□ Stage 1 / Stage 1→2 / Stage 2 / Stage 2→3 체크박스 4개 표시',
        '□ 현재 제품구조의 활성 원료별 체크박스 표시',
        '□ 원료 해제 시 해당 Stage 1 생산지와 Stage 1→2 경로가 지도에서 사라짐',
        '□ Stage 1 해제 시 생산지 원이 사라짐',
        '□ Stage 2→3 해제 시 완성차 경로와 프랑스 시장 마커가 사라짐',
        '□ 6번 분석 탭에서 Scope proxy 그래프 중복이 제거됨',
        '□ CARTO/OSM hosted tile 요청이 없음',
        '□ Natural Earth local polygons가 표시됨',
        '□ OR-Tools 결과값은 지도/UI 변경 전후 동일 입력 기준으로 일치'
    ])

    path=OUT/'EV_Carbon_Optimizer_v0.5.8_전체코드변경_지도라이선스_기술정리.docx'
    doc.save(path)
    print(path)

if __name__=='__main__':
    main()
